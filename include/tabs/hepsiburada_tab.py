from __future__ import annotations

import json
import os
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from PyQt6 import QtCore, QtSvgWidgets, QtWidgets
from PyQt6.QtGui import QCursor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from include.models import APP_TITLE, DEJAVU_FONT_PATH, Product
from include.widgets import ProductTable
from modules.hepsiburada import HepsiburadaScraper


class ProductWorkerSignals(QtCore.QObject):
    item_found = QtCore.pyqtSignal(dict)
    finished = QtCore.pyqtSignal()
    error = QtCore.pyqtSignal(str)


class ProductWorker(QtCore.QRunnable):
    def __init__(
        self,
        scraper: HepsiburadaScraper,
        term: str,
        quick_filters: List[str],
        sorting: Optional[str],
        extra_filters: List[str],
        page_limit: int,
        per_page_limit: int,
    ) -> None:
        super().__init__()
        self.scraper = scraper
        self.term = term
        self.quick_filters = quick_filters
        self.sorting = sorting
        self.extra_filters = extra_filters
        self.page_limit = page_limit
        self.per_page_limit = per_page_limit
        self.sigs = ProductWorkerSignals()

    @QtCore.pyqtSlot()
    def run(self) -> None:
        try:
            products = self.scraper.collect_products(
                self.term,
                self.quick_filters,
                self.sorting,
                self.extra_filters,
                self.page_limit,
                self.per_page_limit,
            )
            for product in products:
                self.sigs.item_found.emit(
                    {
                        "title": product.name,
                        "price": product.price,
                        "link": product.link,
                        "image": product.image,
                        "is_ad": product.is_ad,
                    }
                )
        except Exception as exc:  # pragma: no cover - defensive
            self.sigs.error.emit(str(exc))
        finally:
            self.sigs.finished.emit()


class HepsiburadaTab(QtWidgets.QWidget):
    def __init__(self, parent: Optional[QtWidgets.QWidget] = None):
        super().__init__(parent)
        self.scraper = HepsiburadaScraper()
        self.products: List[Product] = []
        self.dynamic_filters: List[Tuple[str, str]] = []
        self.threadpool = QtCore.QThreadPool.globalInstance()
        self._setup_ui()

    def _setup_ui(self):
        layout = QtWidgets.QVBoxLayout(self)
        layout.addWidget(self._build_header())

        form_layout = QtWidgets.QGridLayout()
        self.search_input = QtWidgets.QLineEdit()
        self.search_input.setPlaceholderText("Aranacak kelime")
        self.search_input.textChanged.connect(self._disable_listing)

        self.page_input = QtWidgets.QSpinBox()
        self.page_input.setMinimum(1)
        self.page_input.setMaximum(50)
        self.page_input.setValue(2)
        self.spinPageCount = self.page_input

        self.per_page_input = QtWidgets.QSpinBox()
        self.per_page_input.setMinimum(1)
        self.per_page_input.setMaximum(200)
        self.per_page_input.setValue(50)
        self.spinProductsPerPage = self.per_page_input

        self.fetch_filters_btn = QtWidgets.QPushButton("Filtreleri Getir")
        self.fetch_filters_btn.clicked.connect(self._load_filters)
        self.list_products_btn = QtWidgets.QPushButton("Ürünleri Listele")
        self.list_products_btn.setEnabled(False)
        self.list_products_btn.clicked.connect(self._list_products)

        form_layout.addWidget(QtWidgets.QLabel("Arama"), 0, 0)
        form_layout.addWidget(self.search_input, 0, 1)
        form_layout.addWidget(QtWidgets.QLabel("Sayfa Sayısı"), 0, 2)
        form_layout.addWidget(self.page_input, 0, 3)
        form_layout.addWidget(QtWidgets.QLabel("Sayfa Başına Ürün"), 0, 4)
        form_layout.addWidget(self.per_page_input, 0, 5)
        form_layout.addWidget(self.fetch_filters_btn, 1, 0)
        form_layout.addWidget(self.list_products_btn, 1, 1)

        layout.addLayout(form_layout)

        filter_layout = QtWidgets.QHBoxLayout()
        filter_layout.addWidget(self._build_quick_filters())
        filter_layout.addWidget(self._build_sorting_filters())
        filter_layout.addWidget(self._build_dynamic_filters())
        layout.addLayout(filter_layout)

        self.table = ProductTable()
        layout.addWidget(self.table)

        action_layout = QtWidgets.QHBoxLayout()
        self.select_all_btn = QtWidgets.QPushButton("Hepsini Seç")
        self.select_all_btn.clicked.connect(lambda: self.table.set_all_checked(True))
        self.clear_selection_btn = QtWidgets.QPushButton("Seçimleri Kaldır")
        self.clear_selection_btn.clicked.connect(lambda: self.table.set_all_checked(False))
        self.remove_selected_btn = QtWidgets.QPushButton("Seçilileri Sil")
        self.remove_selected_btn.clicked.connect(self.table.remove_checked)

        export_menu = QtWidgets.QMenu()
        export_menu.addAction("JSON", self._export_json)
        export_menu.addAction("XLSX", self._export_xlsx)
        export_menu.addAction("PDF", self._export_pdf)
        self.export_btn = QtWidgets.QPushButton("Seçilileri Dışa Aktar")
        self.export_btn.setMenu(export_menu)

        action_layout.addWidget(self.select_all_btn)
        action_layout.addWidget(self.clear_selection_btn)
        action_layout.addWidget(self.remove_selected_btn)
        action_layout.addStretch()
        action_layout.addWidget(self.export_btn)

        layout.addLayout(action_layout)

    def _build_header(self) -> QtWidgets.QWidget:
        widget = QtWidgets.QWidget()
        layout = QtWidgets.QHBoxLayout(widget)
        svg_path = os.path.join("assets", "logo.svg")
        if os.path.exists(svg_path):
            svg_widget = QtSvgWidgets.QSvgWidget(svg_path)
            svg_widget.setFixedSize(96, 96)
            layout.addWidget(svg_widget)
        title_layout = QtWidgets.QVBoxLayout()
        title = QtWidgets.QLabel("Hepsiburada Veri Toplayıcı")
        title.setStyleSheet("font-size: 20px; font-weight: bold;")
        subtitle = QtWidgets.QLabel("Playwright + PyQt6 ile hızlı arama, filtreleme ve dışa aktarma")
        subtitle.setStyleSheet("color: #555;")
        title_layout.addWidget(title)
        title_layout.addWidget(subtitle)
        layout.addLayout(title_layout)
        layout.addStretch()
        status = QtWidgets.QLabel("Durum: Hazır")
        status.setStyleSheet("color: green; font-weight: bold;")
        layout.addWidget(status)
        return widget

    def _build_quick_filters(self) -> QtWidgets.QGroupBox:
        box = QtWidgets.QGroupBox("Hızlı Filtreler")
        layout = QtWidgets.QVBoxLayout(box)
        self.quick_filter_checks: Dict[str, QtWidgets.QCheckBox] = {}
        for label in HepsiburadaScraper.QUICK_FILTERS.keys():
            cb = QtWidgets.QCheckBox(label)
            self.quick_filter_checks[label] = cb
            layout.addWidget(cb)
        return box

    def _build_sorting_filters(self) -> QtWidgets.QGroupBox:
        box = QtWidgets.QGroupBox("Sıralama")
        layout = QtWidgets.QVBoxLayout(box)
        self.sorting_buttons: Dict[str, QtWidgets.QRadioButton] = {}
        for label in HepsiburadaScraper.SORTING_OPTIONS.keys():
            rb = QtWidgets.QRadioButton(label)
            self.sorting_buttons[label] = rb
            layout.addWidget(rb)
        return box

    def _build_dynamic_filters(self) -> QtWidgets.QGroupBox:
        box = QtWidgets.QGroupBox("Kategori Filtreleri")
        layout = QtWidgets.QVBoxLayout(box)
        self.dynamic_filter_container = QtWidgets.QScrollArea()
        self.dynamic_filter_container.setWidgetResizable(True)
        inner = QtWidgets.QWidget()
        self.dynamic_filter_main_layout = QtWidgets.QVBoxLayout(inner)
        self.dynamic_filter_container.setWidget(inner)
        layout.addWidget(self.dynamic_filter_container)
        return box

    def _apply_groupbox_style(self, group_box: QtWidgets.QGroupBox, index: int):
        colors = [
            ("#ff9a00", "#ffb347"),
            ("#ff3d77", "#ff6b9c"),
            ("#00c6ff", "#0072ff"),
            ("#7b2ff7", "#f107a3"),
            ("#42e695", "#3bb2b8"),
            ("#ff512f", "#dd2476"),
        ]
        c1, c2 = colors[index % len(colors)]

        style = f"""
        QGroupBox {{
            margin-top: 12px;
            font-weight: bold;
            border: 2px solid {c2};
            border-radius: 8px;
            padding: 8px;
            background: qlineargradient(
                x1:0, y1:0, x2:1, y2:1,
                stop:0 {c1},
                stop:1 white
            );
        }}
        QGroupBox:title {{
            subcontrol-origin: margin;
            left: 10px;
            top: -2px;
        }}
        """

        group_box.setStyleSheet(style)

    def _styled_slider(self, slider: QtWidgets.QSlider):
        slider.setStyleSheet(
            """
        QSlider::groove:horizontal {
            height: 8px;
            border-radius: 4px;
            background: qlineargradient(
                x1:0, y1:0, x2:1, y2:0,
                stop:0 #ff9a00,
                stop:1 #ff3d77
            );
        }
        QSlider::handle:horizontal {
            background: #ffffff;
            border: 2px solid #ff3d77;
            width: 18px;
            margin: -6px 0;
            border-radius: 9px;
        }
        QSlider::handle:hover {
            background: #ffe5ef;
        }
    """
        )

    def _create_brand_filter_widget(self, items):
        container = QtWidgets.QWidget()
        layout = QtWidgets.QVBoxLayout(container)

        search_box = QtWidgets.QLineEdit()
        search_box.setPlaceholderText("Marka ara…")
        layout.addWidget(search_box)

        list_widget = QtWidgets.QListWidget()
        list_widget.setSelectionMode(QtWidgets.QAbstractItemView.SelectionMode.NoSelection)

        model = QtCore.QStringListModel([name for name, _ in items])
        proxy = QtCore.QSortFilterProxyModel()
        proxy.setSourceModel(model)
        proxy.setFilterCaseSensitivity(QtCore.Qt.CaseSensitivity.CaseInsensitive)

        search_box.textChanged.connect(proxy.setFilterFixedString)

        def refresh_list():
            list_widget.clear()
            for i in range(proxy.rowCount()):
                text = proxy.index(i, 0).data()
                checkbox = QtWidgets.QCheckBox(text)

                for name, query in items:
                    if name == text:
                        checkbox.setProperty("query", query)
                        self.dynamic_filter_checks[name] = checkbox

                item = QtWidgets.QListWidgetItem(list_widget)
                list_widget.setItemWidget(item, checkbox)

        refresh_list()
        proxy.rowsInserted.connect(refresh_list)
        proxy.rowsRemoved.connect(refresh_list)

        layout.addWidget(list_widget)

        return container

    def _render_filters_grouped(self, filters):
        for i in reversed(range(self.dynamic_filter_main_layout.count())):
            item = self.dynamic_filter_main_layout.itemAt(i)
            widget = item.widget()
            if widget:
                widget.setParent(None)

        if isinstance(filters, list):
            grouped = {}
            for name, query in filters:
                if ":" in name:
                    group, item_text = name.split(":", 1)
                else:
                    group, item_text = "Diğer", name
                grouped.setdefault(group.strip(), []).append(
                    {"type": "checkbox", "label": item_text.strip(), "value": query}
                )
            filters = grouped

        self.dynamic_filter_checks = {}

        for idx, (group_name, items) in enumerate(filters.items()):
            group_box = QtWidgets.QGroupBox(group_name)
            group_layout = QtWidgets.QVBoxLayout(group_box)

            for item in items:
                item_type = item.get("type", "checkbox") if isinstance(item, dict) else "checkbox"

                if item_type in {"range", "range-slider"}:
                    h = QtWidgets.QHBoxLayout()
                    self.price_min = QtWidgets.QSlider(QtCore.Qt.Orientation.Horizontal)
                    self.price_max = QtWidgets.QSlider(QtCore.Qt.Orientation.Horizontal)

                    try:
                        min_val = int(item.get("min", 0))
                        max_val = int(item.get("max", 20000))
                    except Exception:
                        min_val, max_val = 0, 20000

                    self.price_min.setRange(min_val, max_val)
                    self.price_max.setRange(min_val, max_val)
                    self.price_min.setValue(min_val)
                    self.price_max.setValue(max_val)

                    self._styled_slider(self.price_min)
                    self._styled_slider(self.price_max)

                    h.addWidget(QtWidgets.QLabel("Min"))
                    h.addWidget(self.price_min)
                    h.addWidget(QtWidgets.QLabel("Max"))
                    h.addWidget(self.price_max)

                    group_layout.addLayout(h)
                    continue

                label_text = item.get("label") if isinstance(item, dict) else str(item)
                query_val = ""
                if isinstance(item, dict):
                    query_val = item.get("query") or item.get("value", "")
                if query_val and not query_val.startswith(("filtreler=", "puan=", "siralama=")):
                    query_val = f"filtreler={query_val}"

                cb = QtWidgets.QCheckBox(label_text)
                cb.setProperty("query", query_val)
                group_layout.addWidget(cb)
                self.dynamic_filter_checks[label_text] = cb

            self._apply_groupbox_style(group_box, idx)
            self.dynamic_filter_main_layout.addWidget(group_box)

        self.dynamic_filter_main_layout.addStretch()

    def _load_filters(self):
        term = self.search_input.text().strip()
        if not term:
            QtWidgets.QMessageBox.warning(self, APP_TITLE, "Lütfen arama kelimesi girin.")
            return
        self.fetch_filters_btn.setEnabled(False)
        QtWidgets.QApplication.setOverrideCursor(QCursor(QtCore.Qt.CursorShape.BusyCursor))
        try:
            filters = self.scraper.fetch_filters(term)
        finally:
            QtWidgets.QApplication.restoreOverrideCursor()
            self.fetch_filters_btn.setEnabled(True)
        if not filters:
            QtWidgets.QMessageBox.warning(self, APP_TITLE, "Filtreler alınamadı, lütfen tekrar deneyin.")
            self.list_products_btn.setEnabled(False)
            return
        self.dynamic_filters = filters
        self._render_filters_grouped(filters)
        self.list_products_btn.setEnabled(True)

    def _collect_quick_filters(self) -> List[str]:
        return [label for label, cb in self.quick_filter_checks.items() if cb.isChecked()]

    def _collect_sorting(self) -> Optional[str]:
        for label, rb in self.sorting_buttons.items():
            if rb.isChecked():
                return label
        return None

    def _collect_dynamic_filters(self) -> List[str]:
        filters: List[str] = []
        if hasattr(self, "dynamic_filter_checks"):
            for _, cb in self.dynamic_filter_checks.items():
                if cb.isChecked():
                    query = cb.property("query")
                    if query:
                        filters.append(query)
        return filters

    def _list_products(self):
        term = self.search_input.text().strip()
        if not term:
            QtWidgets.QMessageBox.warning(self, APP_TITLE, "Lütfen arama kelimesi girin.")
            return
        quick_filters = self._collect_quick_filters()
        sorting = self._collect_sorting()
        extra_filters = self._collect_dynamic_filters()
        page_limit = self.spinPageCount.value()
        per_page_limit = self.spinProductsPerPage.value()
        self.list_products_btn.setEnabled(False)
        self.table.setRowCount(0)
        self.products = []

        worker = ProductWorker(
            self.scraper,
            term,
            quick_filters,
            sorting,
            extra_filters,
            page_limit,
            per_page_limit,
        )
        worker.sigs.item_found.connect(self.add_product_to_table)
        worker.sigs.finished.connect(lambda: self.list_products_btn.setEnabled(True))
        worker.sigs.error.connect(
            lambda msg: QtWidgets.QMessageBox.critical(
                self, APP_TITLE, f"Ürünler alınırken hata oluştu:\n{msg}"
            )
        )
        self.threadpool.start(worker)

    def add_product_to_table(self, item: dict):
        # Temporarily disable sorting while inserting rows to avoid reordering issues
        self.table.setSortingEnabled(False)

        product = Product(
            name=item.get("title", ""),
            price=item.get("price"),
            link=item.get("link", ""),
            image=item.get("image", ""),
            is_ad=item.get("is_ad", False),
        )
        self.products.append(product)
        self.scraper.logger.info("Tabloya eklendi: %s", product.name)

        row = self.table.rowCount()
        self.table.insertRow(row)

        checkbox_item = QtWidgets.QTableWidgetItem()
        checkbox_item.setFlags(checkbox_item.flags() | QtCore.Qt.ItemFlag.ItemIsUserCheckable)
        checkbox_item.setCheckState(QtCore.Qt.CheckState.Unchecked)
        checkbox_item.setData(QtCore.Qt.ItemDataRole.UserRole, product)
        self.table.setItem(row, 0, checkbox_item)

        price_val = item.get("price")
        price_text = "" if price_val is None else str(price_val)

        self.table.setItem(row, 1, QtWidgets.QTableWidgetItem(item.get("title", "")))
        price_item = QtWidgets.QTableWidgetItem(price_text)
        if isinstance(price_val, (int, float)):
            price_item.setData(QtCore.Qt.ItemDataRole.UserRole, float(price_val))
        else:
            price_item.setData(QtCore.Qt.ItemDataRole.UserRole, None)
        self.table.setItem(row, 2, price_item)
        self.table.setItem(row, 3, QtWidgets.QTableWidgetItem(item.get("link", "")))
        self.table.setItem(row, 4, QtWidgets.QTableWidgetItem(item.get("image", "")))
        ad_text = "Evet" if item.get("is_ad", False) else "Hayır"
        self.table.setItem(row, 5, QtWidgets.QTableWidgetItem(ad_text))

        # Re-enable sorting after inserting the row
        self.table.setSortingEnabled(True)

    def _disable_listing(self):
        self.list_products_btn.setEnabled(False)

    def _export_json(self):
        selected = self.table.checked_products()
        if not selected:
            QtWidgets.QMessageBox.information(self, APP_TITLE, "Lütfen dışa aktarmak için ürün seçin.")
            return
        path, _ = QtWidgets.QFileDialog.getSaveFileName(self, "JSON Olarak Kaydet", filter="JSON (*.json)")
        if not path:
            return
        payload = []
        for item in selected:
            payload.append(
                {
                    "name": item.get("name", ""),
                    "price": item.get("price", ""),
                    "link": item.get("link", ""),
                    "image": item.get("image", ""),
                    "is_ad": item.get("is_ad", False),
                }
            )
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        QtWidgets.QMessageBox.information(self, APP_TITLE, "JSON dışa aktarımı tamamlandı.")

    def _export_xlsx(self):
        selected = self.table.checked_products()
        if not selected:
            QtWidgets.QMessageBox.information(self, APP_TITLE, "Lütfen dışa aktarmak için ürün seçin.")
            return
        path, _ = QtWidgets.QFileDialog.getSaveFileName(self, "Excel Olarak Kaydet", filter="Excel (*.xlsx)")
        if not path:
            return
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        import os

        wb = Workbook()
        ws = wb.active
        ws.title = "Hepsiburada"

        headers = ["Ürün Adı", "Fiyat", "Link", "Resim", "Reklam?"]
        ws.append(headers)

        header_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        header_font = Font(bold=True, color="000000")
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        font = Font(name="DejaVu Sans") if os.path.exists(DEJAVU_FONT_PATH) else Font(name="Calibri")

        thin = Side(border_style="thin", color="AAAAAA")
        border = Border(top=thin, left=thin, right=thin, bottom=thin)

        row_index = 2
        for item in selected:
            price_str = item.get("price", "")
            ws.append(
                [
                    item.get("name", ""),
                    price_str,
                    item.get("link", ""),
                    item.get("image", ""),
                    "Evet" if item.get("is_ad", False) else "Hayır",
                ]
            )

            for col in range(1, 6):
                cell = ws.cell(row=row_index, column=col)
                cell.font = font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border

            ad_cell = ws.cell(row=row_index, column=5)
            if ad_cell.value == "Evet":
                ad_cell.fill = PatternFill(start_color="FF9999", fill_type="solid")
            else:
                ad_cell.fill = PatternFill(start_color="CCFFCC", fill_type="solid")

            link_cell = ws.cell(row=row_index, column=3)
            if link_cell.value and link_cell.value.startswith("http"):
                link_cell.value = f'=HYPERLINK("{link_cell.value}", "Link")'
                link_cell.font = Font(color="0000EE", underline="single")

            row_index += 1

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    length = len(str(cell.value))
                    if length > max_len:
                        max_len = length
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 3, 55)

        wb.save(path)
        QtWidgets.QMessageBox.information(self, APP_TITLE, "XLSX dışa aktarımı tamamlandı.")

    def _export_pdf(self):
        selected = self.table.checked_products()
        if not selected:
            QtWidgets.QMessageBox.information(self, APP_TITLE, "Lütfen dışa aktarmak için ürün seçin.")
            return

        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "PDF Olarak Kaydet", filter="PDF (*.pdf)"
        )
        if not path:
            return

        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.utils import ImageReader

        try:
            pdfmetrics.registerFont(TTFont("DejaVu", DEJAVU_FONT_PATH))
            font_name = "DejaVu"
        except Exception:
            font_name = "Helvetica"

        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="NormalTR", fontName=font_name, fontSize=10, leading=13))

        doc = SimpleDocTemplate(
            path,
            pagesize=A4,
            leftMargin=30,
            rightMargin=30,
            topMargin=40,
            bottomMargin=40,
        )

        story = []

        logo_path = "assets/logo.png"
        if os.path.exists(logo_path):
            try:
                story.append(Image(logo_path, width=80, height=80))
                story.append(Spacer(1, 20))
            except Exception:
                pass

        for item in selected:
            name = item.get("name", "")
            price = item.get("price", "")
            link = item.get("link", "")
            image_url = item.get("image", "")
            image_url = self.scraper.upscale_image(image_url)
            if " " in image_url:
                image_url = image_url.split(" ")[0]
            is_ad = "Evet" if item.get("is_ad", False) else "Hayır"

            img_obj = ""
            if image_url.startswith("http"):
                try:
                    img_obj = Image(ImageReader(image_url), width=110, height=110)
                except Exception:
                    img_obj = ""

            card_data = [
                [
                    img_obj,
                    Paragraph(
                        f"<b>{name}</b><br/><br/>"
                        f"<b>Fiyat:</b> {price}<br/>"
                        f"<b>Reklam:</b> {is_ad}<br/><br/>"
                        f"<link href=\"{link}\">{link}</link>",
                        styles["NormalTR"],
                    ),
                ]
            ]

            card = Table(card_data, colWidths=[120, 360], rowHeights=[120])
            card.setStyle(
                TableStyle(
                    [
                        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CCCCCC")),
                        ("BACKGROUND", (0, 0), (-1, -1), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 10),
                        ("TOPPADDING", (0, 0), (-1, -1), 10),
                    ]
                )
            )

            story.append(card)
            story.append(Spacer(1, 20))

        doc.build(story)

        QtWidgets.QMessageBox.information(self, APP_TITLE, "PDF dışa aktarımı tamamlandı.")
    